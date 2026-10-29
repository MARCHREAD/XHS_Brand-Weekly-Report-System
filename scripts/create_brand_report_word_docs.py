from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = "173B70"
BLUE = "2D6CDF"
AMBER = "D99024"
GREEN = "2A9568"
RED = "D14B5A"
INK = "15243C"
MUTED = "64748B"
BORDER = "D9E0EA"
PALE_BLUE = "EFF5FF"
PALE_GRAY = "F7F9FC"
FONT = "Microsoft YaHei"


def font(path: str, size: int, bold: bool = False):
    try:
        return ImageFont.truetype(path, size=size)
    except OSError:
        return ImageFont.load_default()


FONT_PATH = r"C:\Windows\Fonts\msyh.ttc"


def text_center(draw: ImageDraw.ImageDraw, xy, text, f, fill):
    box = draw.multiline_textbbox((0, 0), text, font=f, align="center", spacing=5)
    x, y = xy
    draw.multiline_text((x - (box[2] - box[0]) / 2, y - (box[3] - box[1]) / 2), text, font=f, fill=fill, align="center", spacing=5)


def arrow(draw, start, end, color=BLUE, width=5):
    if not color.startswith("#"):
        color = "#" + color
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        pts = [(x2, y2), (x2 - direction * 17, y2 - 9), (x2 - direction * 17, y2 + 9)]
    else:
        direction = 1 if y2 >= y1 else -1
        pts = [(x2, y2), (x2 - 9, y2 - direction * 17), (x2 + 9, y2 - direction * 17)]
    draw.polygon(pts, fill=color)


def rounded_box(draw, box, title, body, fill, outline, title_color="FFFFFF", body_color="FFFFFF", title_size=27, body_size=19):
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = box
    text_center(draw, ((x1 + x2) / 2, y1 + 42), title, font(FONT_PATH, title_size, True), title_color)
    lines = "\n".join(wrap(body, max(12, int((x2 - x1) / (body_size * 0.95)))))
    text_center(draw, ((x1 + x2) / 2, (y1 + y2) / 2 + 16), lines, font(FONT_PATH, body_size), body_color)


def make_architecture(path: Path):
    im = Image.new("RGB", (1800, 1050), "#F7F9FC")
    d = ImageDraw.Draw(im)
    title_f = font(FONT_PATH, 34, True)
    sub_f = font(FONT_PATH, 20)
    d.text((70, 43), "小红书品牌营销可视化周报系统框架", font=title_f, fill="#15243C")
    d.text((70, 91), "从公开数据采集到周报决策的分层架构", font=sub_f, fill="#64748B")

    rounded_box(d, (70, 180, 380, 420), "数据来源", "小红书公开笔记\n作者资料与评论\n互动指标与话题", "#FFFFFF", "#B9CBE4", "#173B70", "#53627A")
    rounded_box(d, (470, 180, 780, 420), "采集基座", "搜索 / 详情 / 作者 / 评论\n后台任务与请求轨迹\n会话与数据质量状态", "#173B70", "#173B70", "#FFFFFF", "#E5F0FF")
    rounded_box(d, (870, 180, 1180, 420), "数据层", "原始 JSONL / SQLite 快照\n标准化字段与去重\n按周保存可重放数据", "#FFFFFF", "#B9CBE4", "#173B70", "#53627A")
    rounded_box(d, (1270, 180, 1730, 420), "分析层", "声量与互动\n热点与趋势\n情感与评价\n竞品对标与风险", "#2D6CDF", "#2D6CDF", "#FFFFFF", "#EAF2FF")
    arrow(d, (380, 300), (470, 300))
    arrow(d, (780, 300), (870, 300))
    arrow(d, (1180, 300), (1270, 300))

    rounded_box(d, (180, 640, 570, 875), "配置中心", "品牌词、竞品词、行业词\n时间窗与风险阈值", "#FFF5DF", "#E8C27F", "#8A5C14", "#765B32")
    rounded_box(d, (660, 640, 1050, 875), "周报数据集", "BrandWeeklySnapshot\n指标、证据与生成时间\n支持历史周次对比", "#E7F7EF", "#9DD5B9", "#216D4C", "#3B6B57")
    rounded_box(d, (1140, 640, 1530, 875), "门户看板", "总览、热点、竞品\n情绪、行动与归档", "#FFFFFF", "#B9CBE4", "#173B70", "#53627A")
    rounded_box(d, (1560, 640, 1730, 875), "周报导出", "HTML\nPDF\nJSON", "#FFF0F2", "#E6AAB2", "#A63749", "#85505A", 24, 18)
    arrow(d, (1450, 420), (850, 640), color=GREEN)
    arrow(d, (850, 875), (1340, 875), color=GREEN)
    arrow(d, (1530, 755), (1560, 755), color=RED)
    arrow(d, (375, 640), (1400, 420), color=AMBER)

    d.rounded_rectangle((70, 945, 1730, 995), radius=14, fill="#EAF2FF", outline="#C8D9F4", width=2)
    d.text((95, 960), "决策输出：选题规划 · 投放调整 · 舆情回应 · 竞品复盘", font=font(FONT_PATH, 22, True), fill="#24539A")
    im.save(path)


def make_functional(path: Path):
    im = Image.new("RGB", (1800, 1120), "#F7F9FC")
    d = ImageDraw.Draw(im)
    d.text((70, 42), "品牌营销周报门户功能结构", font=font(FONT_PATH, 34, True), fill="#15243C")
    d.text((70, 90), "以周报总览为入口，向热点、竞品、情绪和行动闭环展开", font=font(FONT_PATH, 20), fill="#64748B")

    rounded_box(d, (655, 165, 1145, 305), "品牌营销周报门户", "周选择器 · 品牌工作区 · 数据质量提示", "#173B70", "#173B70", "#FFFFFF", "#E5F0FF", 30, 21)
    modules = [
        ((75, 425, 405, 720), "周报总览", "核心 KPI\n声量与互动趋势\n本周结论"),
        ((465, 425, 795, 720), "热点雷达", "话题排行\n增长曲线\n选题建议"),
        ((855, 425, 1185, 720), "竞品对标", "声量份额\n互动效率\n差异机会"),
        ((1245, 425, 1575, 720), "情绪与评价", "情感分布\n评价维度\n风险证据"),
        ((270, 800, 600, 1035), "行动中心", "行动项\n负责人\n截止时间"),
        ((690, 800, 1020, 1035), "周报归档", "历史快照\n版本对比\n导出记录"),
        ((1110, 800, 1440, 1035), "报表设置", "品牌词\n竞品词\n阈值与权限"),
    ]
    colors = [("#EAF2FF", "#8FB4F2", "#24539A"), ("#FFF5DF", "#E8C27F", "#8A5C14"), ("#E7F7EF", "#9DD5B9", "#216D4C"), ("#FFF0F2", "#E6AAB2", "#A63749"), ("#F2EDFF", "#B9A8E9", "#6847A3"), ("#EFF4FA", "#B9CBE4", "#38577D"), ("#F2F4F7", "#C7D0DC", "#566579")]
    for (box, title, body), (fill, outline, color) in zip(modules, colors):
        rounded_box(d, box, title, body, fill, outline, color, "#53627A", 25, 19)
        arrow(d, (900, 305), ((box[0] + box[2]) // 2, box[1]), color="#8EA5C7", width=4)
    d.rounded_rectangle((70, 1055, 1575, 1098), radius=12, fill="#EAF2FF", outline="#C8D9F4", width=2)
    d.text((90, 1067), "统一能力：数据刷新 · 证据下钻 · 周期切换 · JSON / PDF 导出 · 键盘导航", font=font(FONT_PATH, 19, True), fill="#24539A")
    im.save(path)


def set_run_font(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, head, bold=True, color="FFFFFF", size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell, NAVY)
        if widths:
            cell.width = Inches(widths[i])
    set_repeat_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8.7)
            set_cell_shading(cells[i], "FFFFFF" if ridx % 2 == 0 else PALE_GRAY)
            set_cell_borders(cells[i])
            if widths:
                cells[i].width = Inches(widths[i])
        set_row_cant_split(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def setup_doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.35
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14 if style_name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=24, bold=True, color="000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run(subtitle)
    set_run_font(r, size=11.5, color=MUTED)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("XHS Brand Weekly Report")
    set_run_font(hr, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("小红书品牌营销可视化周报系统 · 设计与开发文档")
    set_run_font(fr, size=8, color=MUTED)
    return doc


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, size=10.5, bold=True, color=INK)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, size=10.5, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=INK)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED, italic=True)


def build_data_doc(path: Path):
    doc = setup_doc("小红书品牌营销周报系统数据需求说明书", "面向数据采集、分析建模与可视化周报开发的字段和口径说明")
    add_para(doc, "本说明书回答系统需要采集什么数据、如何组织数据、如何计算指标，以及如何判断数据是否足以支撑品牌营销决策。数据入口复用现有小红书采集基座，分析层以周为时间单位形成可追溯的 BrandWeeklySnapshot。")

    doc.add_heading("一 建设范围与数据粒度", level=1)
    add_para(doc, "系统围绕品牌、竞品和行业话题建立统一的数据范围。采集数据保持原始响应与请求轨迹，标准化层再按照笔记、作者、评论、话题和周报快照五种粒度组织数据。未返回的字段保持缺失，并记录缺失原因。")
    make_table(doc, ["数据粒度", "主键建议", "时间字段", "主要用途"], [
        ["笔记", "note_id", "published_at / captured_at", "品牌声量、内容类型、互动效率、证据下钻"],
        ["作者", "author_id", "captured_at", "达人结构、作者画像和竞品合作分析"],
        ["评论", "comment_id", "created_at / captured_at", "情绪分析、评价维度、风险预警"],
        ["话题", "topic_id 或标准化词", "week_start / week_end", "热点排行、增长率、选题建议"],
        ["周报快照", "brand_id + week_start", "generated_at", "门户展示、历史对比、导出和审计"],
    ], [1.05, 1.5, 1.65, 3.0])

    doc.add_heading("二 必需原始数据", level=1)
    add_para(doc, "以下字段构成最小可用采集集。字段分为必需、建议和条件必需三类：必需字段缺失时影响主指标；建议字段缺失时保留缺失状态；条件必需字段仅在对应数据链路启用时要求。")
    make_table(doc, ["数据对象", "字段", "级别", "字段说明", "来源"], [
        ["笔记", "note_id / title / desc", "必需", "笔记唯一标识、标题和正文", "搜索与详情"],
        ["笔记", "published_at / type", "必需", "发布时间、图文或视频类型", "详情"],
        ["笔记", "author_id / author_name", "必需", "作者关联及展示名称", "搜索与详情"],
        ["笔记", "liked_count / collected_count / comment_count / share_count", "必需", "互动指标；缺失时保留原始状态", "详情"],
        ["笔记", "tags / topics / image_count", "建议", "话题、标签和内容形态", "详情"],
        ["作者", "author_id / nickname", "必需", "作者主键和昵称", "作者资料"],
        ["作者", "fans_count / interaction_metrics", "建议", "粉丝及公开互动指标", "作者资料 / 蒲公英"],
        ["评论", "comment_id / note_id / content", "必需", "评论主键、所属笔记和文本", "主评论与子评论"],
        ["评论", "created_at / like_count / author_id", "建议", "评论时间、点赞和评论作者", "评论接口"],
        ["任务", "job_id / keyword / pages / status", "必需", "采集任务配置、状态和页数", "任务运行时"],
        ["任务", "captured_at / response_status / error", "必需", "采集时间、响应状态和错误信息", "请求轨迹"],
    ], [0.78, 1.7, 0.75, 3.0, 1.15])

    doc.add_heading("三 品牌与竞品配置数据", level=1)
    add_para(doc, "品牌词配置是分析结果可解释的前提。系统应允许同一品牌维护主品牌词、产品词、别名、排除词和竞品词，并记录版本与生效时间。配置变化后，历史周报保留当时使用的配置版本。")
    make_table(doc, ["配置项", "示例", "用途", "校验要求"], [
        ["品牌主词", "NEO 咖啡", "识别品牌主体", "至少 1 个；大小写和空格标准化"],
        ["品牌别名", "NEO、尼欧咖啡", "覆盖用户自然表达", "去重后保存；记录添加时间"],
        ["产品词", "秋日拿铁、冷萃", "拆分产品线和内容主题", "关联产品或活动实体"],
        ["竞品词", "SEYA COFFEE、NOWWA", "计算 Share of Voice 和差距", "与品牌词使用相同规则"],
        ["排除词", "招聘、二手转让", "降低非营销内容干扰", "支持精确词和正则规则"],
        ["风险阈值", "负向率、增长率、互动量", "触发风险话题", "保存阈值版本和调整人"],
    ], [1.2, 1.7, 2.5, 2.0])

    doc.add_heading("四 派生指标与计算口径", level=1)
    make_table(doc, ["指标", "计算方式", "展示要求"], [
        ["品牌声量", "时间窗内命中品牌词组的去重笔记数", "显示本周值、环比和样本量"],
        ["总互动量", "点赞 + 收藏 + 评论 + 分享；缺失字段不补 0", "显示原始覆盖率和缺失说明"],
        ["互动效率", "互动量 / 笔记数；有曝光字段时再计算互动率", "避免把互动效率称为曝光互动率"],
        ["正向情绪率", "有效文本中正向文本数 / 有效文本总数", "附有效文本量与低置信量"],
        ["热点增长率", "(本周声量 - 上周声量) / 上周声量", "上周样本不足时显示基线不足"],
        ["竞品声量份额", "品牌声量 / 品牌与竞品声量总和", "使用同时间窗和同去重规则"],
        ["风险分", "负向占比、增长速度和互动量的加权结果", "必须可下钻到证据笔记和评论"],
    ], [1.4, 4.1, 2.0])

    doc.add_heading("五 数据质量与治理要求", level=1)
    for item in [
        "完整性：记录每个字段的原始状态、缺失原因和采集覆盖率，空响应不直接等同于没有内容。",
        "唯一性：笔记、作者和评论按平台主键去重；分页重复和重试响应保留请求轨迹。",
        "一致性：同一周报内的品牌词、时间窗、去重规则和情感版本保持一致。",
        "可追溯性：每个结论绑定 evidence，至少能够回到 note_id、comment_id、captured_at 和规则版本。",
        "可重放性：原始快照按周保存，分析任务可离线重算并产生同一份指标结果。",
        "隐私边界：仅保存完成分析所需的公开字段，不将登录 Cookie、设备配置和私密会话写入周报数据集。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("六 周报输出数据集", level=1)
    add_para(doc, "门户读取统一的 BrandWeeklySnapshot。该对象同时保存指标、趋势、洞察和证据索引，避免前端直接拼接原始响应。")
    make_table(doc, ["字段组", "核心字段", "说明"], [
        ["身份", "brand_id / week_start / week_end", "品牌和报告时间窗"],
        ["质量", "sample_size / coverage / generated_at", "样本量、覆盖率和生成时间"],
        ["KPI", "volume / interactions / positive_rate / risk_topics", "四项总览指标"],
        ["趋势", "period / volume / interactions", "按日或按周的趋势序列"],
        ["主题", "topic / volume / growth / sentiment", "热点及其情绪状态"],
        ["竞品", "brand / share / interactions / gap", "竞品对标和差异机会"],
        ["证据", "note_id / comment_id / matched_terms / url", "结论下钻和审计依据"],
    ], [1.2, 3.0, 3.3])

    doc.add_heading("七 最小可用版本", level=1)
    add_para(doc, "第一期只要稳定获得笔记搜索结果、笔记详情、主评论、发布时间、作者标识和互动指标，就可以上线周报总览、热点排行和基础情绪分析。作者画像、子评论、蒲公英商业数据和自动导出属于第二阶段增强项。")
    doc.add_heading("八 数据接口交付建议", level=1)
    add_para(doc, "分析层与门户之间建议使用版本化 JSON 契约。接口只返回周报快照和必要的证据索引，原始响应通过 note_id 或 comment_id 进行二次查询，避免首屏负载过大。")
    make_table(doc, ["接口", "方法", "用途", "最低返回"], [
        ["/v1/brand-report/snapshots", "GET", "按品牌和周次读取周报", "brand_id、week、kpis、quality"],
        ["/v1/brand-report/topics", "GET", "读取热点排行和证据", "topic、volume、growth、evidence"],
        ["/v1/brand-report/competitors", "GET", "读取竞品对标", "brand、share、interactions、gap"],
        ["/v1/brand-report/sentiment", "GET", "读取情绪和评价维度", "distribution、aspects、examples"],
        ["/v1/brand-report/runtime", "GET", "读取采集新鲜度和质量", "last_sync、job_status、coverage"],
    ], [2.15, 0.8, 2.2, 2.8])
    doc.add_heading("九 交付优先级", level=1)
    for item in [
        "P0：笔记主键、标题、正文、发布时间、作者标识、互动指标、采集时间和任务状态。",
        "P1：主评论、话题标签、品牌与竞品配置、去重规则、BrandWeeklySnapshot 和证据索引。",
        "P2：作者画像、子评论、蒲公英商业字段、低置信文本复核和 HTML/PDF 自动导出。",
    ]:
        add_bullet(doc, item)
    doc.save(path)


def build_system_doc(path: Path, arch: Path, functional: Path):
    doc = setup_doc("小红书品牌营销周报系统架构与功能说明书", "系统框架图、功能结构图及模块职责说明")
    add_para(doc, "本说明书用于说明品牌营销周报系统如何承接现有小红书数据采集基座，并将原始内容转化为可追溯的品牌营销洞察。阅读者可以从两张图快速理解系统边界，再结合文字说明定位开发模块和验收重点。")

    doc.add_heading("一 系统框架图", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(arch), width=Inches(6.85))
    add_caption(doc, "图 1  从小红书公开数据到品牌营销决策的分层框架")
    add_para(doc, "系统采用数据来源、采集基座、数据层、分析层和展示层五层结构。数据来源提供公开笔记、作者、评论和互动信息；采集基座复用现有搜索、详情、作者和评论能力，并负责任务状态与请求轨迹；数据层保存原始快照和标准化结果；分析层计算声量、互动、热点、情绪、竞品和风险；展示层以品牌周报门户和导出文件服务营销决策。")
    add_para(doc, "配置中心为分析层提供品牌词、竞品词、行业词、排除词和风险阈值。数据质量与证据索引贯穿各层，确保门户中的增长、领先和风险结论都可以回到具体样本。")

    doc.add_heading("二 框架分层职责", level=1)
    make_table(doc, ["层级", "核心职责", "输入", "输出"], [
        ["数据来源", "定义公开数据范围和平台字段", "公开笔记、作者、评论", "可采集对象"],
        ["采集基座", "请求、签名、分页、重试和任务管理", "关键词、分页、会话状态", "原始响应与任务轨迹"],
        ["数据层", "保存、去重、标准化和按周快照", "原始响应", "NormalizedNote / Comment"],
        ["分析层", "计算指标、识别热点、评估情绪和竞品差距", "标准化数据、配置规则", "BrandWeeklySnapshot"],
        ["展示层", "门户看板、下钻证据和周报导出", "周报快照、运行状态", "门户、HTML/PDF/JSON"],
    ], [1.15, 2.65, 1.7, 2.0])

    doc.add_heading("三 功能结构图", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(functional), width=Inches(6.85))
    add_caption(doc, "图 2  品牌营销周报门户功能结构")
    add_para(doc, "门户以周报总览为一级入口，核心分析模块围绕四个问题组织：品牌本周表现如何、用户正在讨论什么、竞品在哪些话题上领先、用户为什么喜欢或不满。行动中心把结论转成负责人和截止时间，周报归档保存历史快照，报表设置维护品牌词和分析阈值。")

    doc.add_heading("四 功能模块文字说明", level=1)
    make_table(doc, ["模块", "功能说明", "关键页面元素", "依赖数据"], [
        ["周报总览", "展示本周核心 KPI、环比变化和自动摘要", "声量、互动、情绪、风险卡片；趋势图；本周结论", "BrandWeeklySnapshot.kpis / trend / insights"],
        ["热点雷达", "发现增长最快的话题并提供内容证据", "话题排行、增长率、情绪、样本笔记", "topics / notes / matched_terms"],
        ["竞品对标", "比较品牌与竞品的声量份额和互动效率", "Share of Voice、进度条、机会提示", "competitors / volume / interactions"],
        ["情绪与评价", "拆分正向、中性、负向声音并识别评价维度", "情绪环图、关键词、原声评论、风险证据", "sentiment / comments / evidence"],
        ["行动中心", "把洞察转成可跟踪的营销任务", "行动项、负责人、截止时间、优先级", "insights / action_items"],
        ["周报归档", "按品牌和周次查看历史快照并导出", "周次筛选、版本差异、导出记录", "weekly_snapshots / export_logs"],
        ["报表设置", "维护品牌、竞品、行业词和阈值", "词库版本、阈值、权限和采集范围", "brand_config / rule_versions"],
    ], [1.15, 2.45, 2.25, 1.65])

    doc.add_heading("五 典型使用流程", level=1)
    for item in [
        "配置品牌工作区：录入品牌主词、别名、产品词、排除词和竞品词，并保存配置版本。",
        "执行周采集：采集基座按关键词和任务配置获取笔记、作者、评论及互动数据，记录任务状态和覆盖率。",
        "生成周报快照：标准化层去重和字段映射，分析层计算指标、话题、情绪、竞品差距和风险分。",
        "审核洞察证据：营销人员从总览下钻到笔记和评论，确认结论的样本量、时间窗和证据。",
        "形成行动计划：将热点机会、竞品差异和风险反馈转成内容、投放、产品或客服任务。",
        "归档与复盘：保存周报快照和导出版本，在下一周对比增长、行动完成度和风险变化。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("六 角色与权限建议", level=1)
    make_table(doc, ["角色", "可见范围", "主要操作"], [
        ["品牌负责人", "全部品牌周报和归档", "查看结论、确认行动、导出周报"],
        ["内容运营", "热点、笔记和行动中心", "筛选话题、查看证据、维护选题任务"],
        ["媒介投放", "竞品与达人相关模块", "查看声量份额、互动效率和作者结构"],
        ["公关客服", "情绪、评价和风险证据", "查看负向样本、记录回应和处理状态"],
        ["数据管理员", "配置、任务、质量和审计", "维护词库、阈值、采集任务和版本"],
    ], [1.25, 2.55, 3.7])

    doc.add_heading("七 非功能要求与验收重点", level=1)
    for item in [
        "可用性：门户在 375px、768px、1024px 和 1440px 宽度下保持可读和可操作，不出现横向滚动。",
        "可追溯性：每个 KPI 和洞察都能查看时间窗、样本量、覆盖率、规则版本和证据对象。",
        "可恢复性：采集失败时保留上一版快照，并在门户中显示采集状态和数据新鲜度。",
        "可重算性：同一份离线快照重复执行分析任务，应得到一致的指标和排序结果。",
        "性能：门户首屏只读取周报快照，原始响应和证据在用户下钻时按需加载。",
        "安全：账号会话、Cookie、设备私有配置和运行数据库不进入周报导出文件。",
    ]:
        add_bullet(doc, item)
    doc.save(path)


if __name__ == "__main__":
    arch = ASSETS / "brand-report-system-architecture.png"
    functional = ASSETS / "brand-report-functional-structure.png"
    make_architecture(arch)
    make_functional(functional)
    build_data_doc(DOCS / "brand-weekly-report-data-requirements.docx")
    build_system_doc(DOCS / "brand-weekly-report-system-diagrams.docx", arch, functional)
    print("created")
