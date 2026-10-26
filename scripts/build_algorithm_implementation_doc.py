"""Build the algorithm and implementation supplement using the retained course template.

Use the Codex bundled Python runtime. Screenshots must already exist in the task directory.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

ROOT=Path(__file__).resolve().parents[1]
TMP=ROOT/'output/algorithm-implementation-20261026'
REF=Path('A:/download/智慧城市技术概论 大作业模板.docx')
OUT=ROOT/'docs/李昊阅_202409010201_智能算法原理与实现.docx'
ASSETS=ROOT/'docs/assets/algorithm-implementation'
RUNTIME=ROOT/'platforms/xiaohongshu/output/weekly/7cf01c2524d14b1482ecc9dba5aef72a'


def run():
    digest=hashlib.sha256(REF.read_bytes()).hexdigest()
    evidence=json.loads((RUNTIME/'implementation-evidence.json').read_text(encoding='utf-8'))
    assert evidence['evaluation']['reproduced'] and evidence['evaluation']['bertopic_reload']
    assert evidence['topic']['fit_sample_size']==1058
    ASSETS.mkdir(parents=True,exist_ok=True)
    # Preserve native captures. No compositing, retouching, or invented terminal output.
    for name in ['dev-bertopic.png','dev-svm.png','model-dashboard.png','real-dashboard.png']:
        shutil.copyfile(TMP/name,ASSETS/name)
    (ASSETS/'verified-metrics.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    contract=f'''# Template execution contract
Reference: {REF}; SHA256 {digest}. Rendered reference pages: reference/page-1.png through page-4.png.
Retain A4 portrait, 11906 by 16838 twips, top/bottom margins 1440, left/right 1800, single column.
Retain section properties, header1.xml, footer1.xml, original numbering, styles, font table, theme and opaque parts byte for byte.
Use source Heading 2 and Heading 3 components for numbered headings; source title centered 22pt bold Song and identity 14pt bold.
Use Song 12pt, line spacing 1.25, two-character first-line indent, no empty body paragraphs. Captions centered 10.5pt.
This supplement fills the source algorithm and practical-implementation content patterns; other template instructional placeholders are removed.
Seven pages: model/data scope; BERTopic principle; LinearSVC principle; measured effects; BERTopic IDE; SVM IDE; running dashboard and reproduction.
Add only new document image relationships; keep original relationships and all original assets. Images occupy about 90 percent of text width.
Editable package parts: document.xml, document.xml.rels, settings.xml, core metadata and added PNG images. Append only the missing Title style, preserving all original style nodes.
Use actual saved metrics and a reproduced fixed test set. Training labels include AI-assisted labels; negative recall is zero. No invented topic precision or business lift.
Images are native VS Code captures and live localhost screenshots, with no fabricated interface or terminal output.
Verify final pages visually and compare preserved package parts. Original template must remain unchanged.
'''
    (TMP/'artifact.md').write_text(contract,encoding='utf-8')
    doc=Document(REF)
    if 'Title' not in doc.styles:
        title_style=doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
        title_style.base_style=doc.styles['Normal']
    body=doc._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'):body.remove(child)
    # Existing styles remain the authority; set exact run/paragraph properties locally.
    def p(text,role='body',new_page=False):
        para=doc.add_paragraph(style={'body':'Normal','h1':'Heading 2','h2':'Heading 3','title':'Title','caption':'Normal','name':'Normal'}[role])
        fmt=para.paragraph_format
        fmt.page_break_before=new_page
        fmt.keep_with_next=role in {'h1','h2','title'}
        fmt.keep_together=True
        fmt.space_before=Pt(0 if role=='body' else 5)
        fmt.space_after=Pt(0 if role=='body' else 6)
        fmt.line_spacing=1.25 if role=='body' else 1.0
        fmt.first_line_indent=Pt(24) if role=='body' else Pt(0)
        para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if role=='body' else (WD_ALIGN_PARAGRAPH.CENTER if role in {'title','caption','name'} else WD_ALIGN_PARAGRAPH.LEFT)
        if role not in {'h1','h2'}:
            np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);para._p.get_or_add_pPr().append(np)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');para._p.get_or_add_pPr().append(snap)
        run=para.add_run(text)
        run.font.name='宋体';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        run.font.size=Pt({'body':12,'h1':18,'h2':16,'title':22,'caption':10.5,'name':14}[role])
        run.font.bold=role in {'h1','h2','title','name'};run.font.color.rgb=RGBColor(0,0,0)
        return para

    def math(text):
        para=p('',role='caption');para.paragraph_format.space_before=Pt(5)
        om=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
        def run(parent,value):
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=value;r.append(t);parent.append(r)
        scripts={'eᵢ':('e','i',None),'eⱼ':('e','j',None),'Σᵤ':('Σ','u',None),
                 'Σᵥ':('Σ','v',None),'Σᵢ':('Σ','i',None),'TFIDF₂₋₄':('TFIDF','2:4',None),
                 'αᵧᵢ':('α','y_i',None),'yᵢ':('y','i',None),'xᵢ':('x','i',None),
                 '‖₂²':('‖','2','2'),'‖₂':('‖','2',None),']²':(']',None,'2')}
        pattern='('+ '|'.join(re.escape(token) for token in sorted(scripts,key=len,reverse=True))+')'
        for token in re.split(pattern,text):
            if not token:continue
            if token not in scripts:run(m,token);continue
            base,sub,sup=scripts[token]
            node=OxmlElement('m:sSubSup' if sub and sup else ('m:sSub' if sub else 'm:sSup'))
            el=OxmlElement('m:e');run(el,base);node.append(el)
            for name,value in [('sub',sub),('sup',sup)]:
                if value is not None:el=OxmlElement('m:'+name);run(el,value);node.append(el)
            m.append(node)
        om.append(m);para._p.append(om)

    def table(headers,rows,widths):
        tab=doc.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
        for c,w in zip(tab.columns,widths):c.width=Inches(w)
        for i,values in enumerate([headers]+rows):
            row=tab.rows[0] if i==0 else tab.add_row()
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if i==0:trpr.append(OxmlElement('w:tblHeader'))
            for j,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                pr.append(borders)
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EDF4' if i==0 else ('F7F9FC' if i%2==0 else 'FFFFFF'));pr.append(fill)
                margins=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'65');x.set(qn('w:type'),'dxa');margins.append(x)
                pr.append(margins)
                para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=1
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or j==0 else WD_ALIGN_PARAGRAPH.LEFT
                run=para.add_run(str(value));run.font.name='宋体';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');run.font.size=Pt(10.5);run.font.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)
        return tab

    def fig(name,caption,width=5.2):
        para=p('',role='caption');para.paragraph_format.keep_with_next=True
        run=para.add_run();run.add_picture(str(ASSETS/name),width=Inches(width))
        for pr in para._p.xpath('.//wp:docPr'):pr.set('descr',caption)
        p(caption,'caption')

    # Page 1: scope and the specific two algorithms.
    p('小红书品牌周报系统\n智能算法原理与实现','title')
    p('姓名：李昊阅    学号：202409010201','name')
    p('系统涉及的相关智能算法','h1')
    p('模型选择与处理数据','h2')
    p('本系统采用 BERTopic 进行品牌讨论话题发现，采用支持向量机的具体实现 LinearSVC 进行评论情感三分类。前者把相近讨论聚成话题并提取关键词，后者将评论识别为正向、中性或负向，形成可查看原文证据的营销分析结果。')
    p('研究对象为 libtv，竞品为即梦和 updream。模型使用真实采集的笔记标题、正文及所属评论；品牌归因依据标题或正文中的名称匹配，搜索命中只作为来源记录。评论沿所属讨论串归因，不代表每条评论都在直接评价工具产品。')
    p('数据清洗包括主键去重、链接与表情标记处理、重复文本合并、过短文本和订阅提醒过滤。过滤仅作用于分析输入，已采原始数据保留。周报按真实发布时间筛选；话题管线在已采跨期语料上拟合，本期周报再统计相应日期内的文本。')
    p('表1 两种算法的数据与输出','caption')
    table(['具体模型','处理的数据','输出及用途'],[
        ['BERTopic','品牌相关标题、正文和评论\n1,737 条输入，1,058 条有效去重文本','话题编号、关键词、文本量\n定位创作需求、积分与会员讨论'],
        ['LinearSVC','有正中负标签的真实评论\n253 条有效去重标注','情感类别与分类分数\n汇总评价、筛选待复核评论'],
    ],[1.05,2.15,2.0])
    p('本次模型记录生成于 2026 年 10 月 26 日 21:17，21:24 对保存的模型进行了离线加载复核。复核时已发现 3,850 篇笔记、保存 7,800 条评论，301 篇达到当前采集目标，整体采集仍为部分完成；这些总数与模型实际使用的样本量分别统计。')
    p('当前评论范围为每篇前两页主评论及随页回复，保留调整前已采数据。样本包含早期更深的评论覆盖，所得占比反映已采样本，未作平台总体或销量推断。')

    # Page 2: exact topic pipeline and what each stage computes.
    p('BERTopic 话题发现原理','h2',True)
    p('BERTopic 是一种基于语义向量和密度聚类的话题建模方法。本系统固定采用 sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2 编码器，串联 UMAP、HDBSCAN 与 c-TF-IDF。编码器使用预训练权重，项目拟合的是后续降维、聚类和话题表示管线。')
    p('第一步是语义编码。MiniLM 将每条文本转为 384 维向量并进行 L2 归一化，使表达接近的内容在向量空间中相邻。例如“会员积分过期”和“积分到期清零”可通过语义接近性形成关联。该例用于解释算法，不作为本次模型正确率证据。')
    math('cos(eᵢ, eⱼ) = (eᵢ · eⱼ) / (‖eᵢ‖₂ ‖eⱼ‖₂)')
    p('第二步使用 UMAP 降维。UMAP 先根据高维向量构造近邻图，再优化低维图与高维图的邻接关系，使局部相似文本尽量保持邻近。本次将 384 维降至 5 维，n_neighbors=15、min_dist=0、metric=cosine、random_state=42，为后续密度聚类提供输入。')
    p('第三步使用 HDBSCAN 聚类。算法根据核心距离与互达距离建立密度层次，选择相对稳定的密集簇；低密度、难归属的文本标为 -1。它无需事先指定话题数量。本次 min_cluster_size=30、min_samples=3，使用欧氏距离；6 个话题是拟合结果。')
    p('第四步使用 c-TF-IDF 提取关键词。先将同一簇中的文本合并为一个类别文档，再计算词在该类中的相对频率与跨类别区分能力。高分词构成话题的解释，例如“积分、会员、到期”。中文采用 jieba 分词，过滤停用词、无信息符号及品牌词，避免只按品牌名形成话题。')
    math('w(t,c) = (f(t,c) / Σᵤ f(u,c)) × ln(1 + A / Σᵥ f(t,v))')
    p('式中 f(t,c) 为词 t 在话题 c 中的频数，A 为各话题合并文档的平均总词数。该式描述本系统默认 c-TF-IDF 的核心加权形式。每个话题保留排名靠前的关键词与关联文本量，供运营人员查看和人工命名。')
    p('该算法增加的智慧性在于：从大量非结构化讨论中自动归纳相近主题，并允许噪声文本保留未归类状态，减少逐条阅读。周报再按品牌和周次统计话题分布；跨次重训的话题编号变化，需先对齐含义再比较趋势。')

    # Page 3: TF-IDF, frozen vectors, squared-hinge LinearSVC and group evaluation.
    p('LinearSVC 情感分类原理','h2',True)
    p('本系统采用 scikit-learn 的 LinearSVC，属于线性支持向量机。输入是带标签的单条评论，输出是 negative、neutral、positive 三类。特征由字符 TF-IDF 和固定 MiniLM 语义向量组合，既保留“好用”“太贵”等局部表达，也引入句子语义信息。')
    p('字符 TF-IDF 按连续字符片段提取特征。本次选中的 n-gram 范围为 2 至 4，最高保留 60,000 个特征，并启用 sublinear_tf。词频部分使用 1+ln(tf)，逆文档频率使用平滑形式 ln((1+N)/(1+df))+1，再作 L2 归一化。TF-IDF 只在训练数据上拟合，测试数据不参与词表和权重学习。')
    math('x = [ TFIDF₂₋₄(text) ; 0.5 × MiniLM(text) ]')
    p('其中 MiniLM 使用固定预训练权重，不参与有监督微调。两组特征拼接后进入线性分类器。SVM 寻找分类超平面 w·x+b，使不同类别尽量分离，同时控制模型复杂度；对难分样本允许一定分类损失，以适应真实评论中的混合表达。')
    math('min  ½‖w‖₂² + C Σᵢ αᵧᵢ [max(0, 1 − yᵢ(w·xᵢ+b))]²')
    p('上式为二分类子问题的 L2 正则化平方铰链损失形式。LinearSVC 采用一对其余方式训练三个分类器；α 为 balanced 类别权重，C 调节间隔和误分损失的权衡。本次 C=1.0、max_iter=10000、random_state=42，参数由训练内部三折分组交叉验证选定。')
    p('预测时计算三个类别的 decision_function 分数，取分数最高的类别。若最高和次高分数差小于 0.2，则转为“待复核”。这个分差用于识别边界模糊的样本，不解释为概率，也不是经校准的置信度。')
    p('训练与评测按笔记分组隔离，固定保留测试评论及其所属整篇笔记，防止同一讨论串同时进入训练和测试。本次 205 条用于评估模型训练，29 条用于固定测试，19 条来自保留笔记但不计入测试。评估完成后另行用全部 253 条标签重拟合生产模型；评估模型单独保存，因此页面分数来自固定测试流程。')
    p('该算法增加的智慧性在于：对新评论执行自动情感分类，减少人工初筛工作，并把低分差样本交给人工复核。实际负向识别效果仍需根据下一页评测结果判断。')

    # Page 4: measured effect, full denominator and limits.
    p('真实数据上的使用效果','h2',True)
    p('BERTopic 共处理 1,737 条品牌相关输入，清洗与去重后保留 1,058 条，其中 959 条进入 6 个话题，99 条标为噪声，占 9.36%。这些数值衡量分组覆盖情况；当前未建立人工话题真值集，尚未报告话题准确率或营销效果提升比例。')
    p('表2 BERTopic 的实际话题规模','caption')
    table(['话题','关键词概括','文本数'],[
        ['0','更新、评论、好看等泛互动讨论','659'],
        ['1','AI 生成、创作、工具与漫剧制作','100'],
        ['2','积分、会员、到期与使用','71'],
        ['3','AI 创作者活动与成长计划','48'],
        ['4','积分到期、重置与交易讨论','48'],
        ['5','账号称呼及零散互动','33'],
    ],[.6,3.8,.8])
    p('表中名称由关键词概括。泛互动话题占有效文本的 62.29%，说明模型仍受到作品讨论与聊天噪声影响；6 个话题不等于 6 个有效营销热点。积分与会员相关话题可作为价格机制、使用门槛等运营问题的复核线索。')
    p('表3 固定测试集上的 LinearSVC 评测','caption')
    table(['指标','实测结果','对应样本或含义'],[
        ['准确率','89.66%','29 条中正确 26 条'],
        ['宏平均 F1','0.6050','三类 F1 等权平均'],
        ['负向召回率','0.00%','2 条负向均被判为中性'],
        ['中性召回率','100.00%','8 条中性均被识别'],
        ['正向召回率','94.74%','19 条中识别 18 条'],
    ],[1.25,1.0,2.95])
    p('标签共 253 条，其中人工复核 86 条、AI 辅助标注 167 条；固定测试集的 29 条全部为辅助标签。这一成绩反映对当前标签集的拟合一致性，独立人工评测仍待补充。负向样本识别失效，现阶段风险提示应作为人工复核候选；优先补充不同笔记中的真实负向评论和反讽、混合评价样本。')
    p('三折分组交叉验证的宏平均 F1 为 0.5254；前后两轮在相同 29 条测试评论上的宏平均 F1 均为 0.6050，当前没有观察到测试指标提升。系统已实现话题聚合、情感初筛和证据下钻，实际节省工时及营销收益尚待用户试用测量。')

    # Page 5: native development evidence for BERTopic.
    p('实践','h1',True)
    p('开发环境与 BERTopic 实现','h2')
    p('开发环境为 Windows、VS Code、Python 3.13 及平台独立虚拟环境 .venv-weekly，前端由本机 FastAPI 服务提供。主要依赖为 BERTopic 0.17.4、sentence-transformers 5.7.0、umap-learn 0.5.12、hdbscan 0.8.44 和 scikit-learn 1.7.2，模型在 CPU 上运行。')
    fig('dev-bertopic.png','图1 VS Code 中的 BERTopic 管线代码与实际模型复核终端')
    p('具体实现位于 collector/weekly/topics.py。train_topics 依次构造语义向量、UMAP、HDBSCAN 和中文 CountVectorizer，调用 fit_transform 获得话题归属，使用 c-TF-IDF 排序话题词，最后输出话题关键词和各簇文本数。')
    p('完整拟合管线保存为 bertopic-model.joblib，固定编码器按模型 ID 重新挂接。复核脚本实际加载该文件，并对三条已有向量执行 transform，得到的话题编号与保存记录一致；这验证了模型可加载并继续预测。')
    p('图1为本机 VS Code 实际窗口截图。复核脚本重新计算固定测试结果，并检查保存模型的 SHA-256；模型本体、原始评论及会话信息均保留在本地运行目录。')
    p('实现过程将话题预测和周报统计分开：模型在跨期文本中学习主题结构，周报接口再以发布时间筛选本期语料。这样既能使用更多已采文本，也能避免把历史讨论量直接作为本周热点声量。')

    # Page 6: native SVM source and reproducible confusion matrix.
    p('LinearSVC 训练与结果复现','h2',True)
    fig('dev-svm.png','图2 VS Code 中的分组交叉验证代码和固定测试集运行结果')
    p('训练实现位于 collector/weekly/sentiment.py，语义特征封装位于 semantic.py。Pipeline 将特征构造与 LinearSVC 串联；GridSearchCV 只在训练部分比较字符片段范围和 C 值，以三折分组交叉验证的宏平均 F1 选择参数。')
    p('截图终端显示评估模型训练 205 条、固定测试 29 条，以及重新得到的 Accuracy=89.6552%、Macro-F1=0.6050、Negative recall=0.0000。脚本还加载生产分类器进行预测，固定测试复现与模型重载检查均通过。')
    p('表4 复现得到的混淆矩阵','caption')
    table(['真实类别／预测类别','负向','中性','正向'],[
        ['负向','0','2','0'],['中性','0','8','0'],['正向','0','1','18'],
    ],[2.2,1.0,1.0,1.0])
    p('矩阵按行是真实标签、列是模型预测排列。两个负向样本流入中性列，是目前最明确的误差来源。正向类样本较多，使总体准确率较高，因此系统同时展示宏平均 F1、每类召回率及测试样本量。')
    p('本次复核执行于 2026 年 10 月 26 日 21:24，仅加载模型和本地数据，不触发采集，也不重新训练。测试结果与 21:17 保存的评估记录一致。')

    # Page 7: user-facing output and concrete reproduction.
    p('模型结果在系统中的应用','h2',True)
    fig('model-dashboard.png','图3 本地真实周报的模型样本量 指标与混淆矩阵')
    p('本地页面 http://127.0.0.1:8000/brand-report/ 已接入真实模型记录。话题结果用于组织关注点，SVM 预测用于情绪分布及待复核评论；页面同步显示标签来源、样本覆盖、各类召回率和保存模型的加载状态。图3 的模型指标对应本次保存的 21:17 分析版本。')
    p('可复核的运行方法','h2')
    p('在 platforms/xiaohongshu 目录打开 VS Code 终端，使用 .venv-weekly 的 Python 运行 scripts/verify_weekly_models.py，并传入已有任务 ID。该脚本复算固定测试集和混淆矩阵，检查 BERTopic 加载预测及模型哈希，最后输出复核 JSON。')
    cmd=p('.\\.venv-weekly\\Scripts\\python.exe scripts\\verify_weekly_models.py\n--job 7cf01c2524d14b1482ecc9dba5aef72a','caption')
    cmd.paragraph_format.keep_together=True
    p('上方命令在终端实际输入时合为一行。查看应用时运行同目录 main.py，再访问周报入口；更新分析使用 weekly.py analyze --job 任务ID，增加标注后使用 weekly.py train --job 任务ID。复现检查与训练操作分别执行，以保留固定评测依据。')
    p('本阶段完成了具体算法落地、真实数据处理、模型持久化、固定测试复现及页面集成。下一步围绕人工负向标签、泛互动过滤和跨次话题对齐完善结果，再用新的独立人工测试集评估泛化表现。')

    settings=doc.settings.element
    update=settings.find(qn('w:updateFields'))
    if update is None:update=OxmlElement('w:updateFields');settings.append(update)
    update.set(qn('w:val'),'true')
    doc.core_properties.title='小红书品牌周报系统智能算法原理与实现'
    doc.core_properties.author='李昊阅'
    authored=BytesIO();doc.save(authored)
    with ZipFile(REF) as original,ZipFile(BytesIO(authored.getvalue())) as generated:
        original_parts={name:original.read(name) for name in original.namelist()}
        generated_parts={name:generated.read(name) for name in generated.namelist()}
    allowed={'word/document.xml','word/_rels/document.xml.rels','word/settings.xml','docProps/core.xml','word/styles.xml'}
    final=dict(original_parts)
    for name,data in generated_parts.items():
        if name in allowed or name not in original_parts:final[name]=data
    from lxml import etree
    original_styles=etree.fromstring(original_parts['word/styles.xml'])
    generated_styles=etree.fromstring(generated_parts['word/styles.xml'])
    original_ids={node.get(qn('w:styleId')) for node in original_styles.findall(qn('w:style'))}
    for node in generated_styles.findall(qn('w:style')):
        if node.get(qn('w:styleId')) not in original_ids:
            assert node.get(qn('w:styleId'))=='Title'
            original_styles.append(deepcopy(node))
    final['word/styles.xml']=etree.tostring(original_styles,encoding='UTF-8',xml_declaration=True,standalone=True)
    with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
        for name,data in final.items():z.writestr(name,data)
    assert hashlib.sha256(REF.read_bytes()).hexdigest()==digest
    assert all(final[name]==data for name,data in original_parts.items() if name not in allowed)
    manifest={'reference_sha256':digest,'output':str(OUT),'preserved_original_parts':len(original_parts)-len(allowed),
              'changed_original_parts':[name for name in original_parts if final[name]!=original_parts[name]],
              'screenshots':{name:hashlib.sha256((ASSETS/name).read_bytes()).hexdigest() for name in ['dev-bertopic.png','dev-svm.png','model-dashboard.png','real-dashboard.png']}}
    (TMP/'verification.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(OUT)


if __name__=='__main__':run()
